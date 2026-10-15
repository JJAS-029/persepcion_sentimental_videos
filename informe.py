"""Informe de comentarios en el formato solicitado (un bloque por video / plataforma):

    <Página>
     <Plataforma> <link>

    🟢<Autor> <comentario positivo>
    🔴<Autor> <comentario negativo>
"""
import io
import re

import pandas as pd

ICONOS = {"POS": "🟢", "NEG": "🔴", "NEU": "⚪"}


def seleccionar(df: pd.DataFrame, por_tipo: int | None, incluir_neutros: bool = False) -> pd.DataFrame:
    """Comentarios para el informe, agrupados por tipo y en orden de hora.

    por_tipo=N: los N más claros de cada tipo (mayor probabilidad), sin repetidos ni mensajes de puro emoji.
    por_tipo=None: todos los comentarios de cada tipo, sin filtrar.
    """
    tipos = [("POS", "p_pos"), ("NEG", "p_neg")] + ([("NEU", "p_neu")] if incluir_neutros else [])
    if por_tipo is None:
        return pd.concat([df[df["sentimiento"] == s].sort_values("hora") for s, _ in tipos])
    d = df.copy()
    d["_norm"] = d["mensaje"].str.lower().str.replace(r"\s+", " ", regex=True).str.strip()
    d = d[d["_norm"].str.count(r"[a-záéíóúñü]") >= 8]          # descarta solo emojis / "hola"
    d = d.drop_duplicates("_norm")
    bloques = [d[d["sentimiento"] == s].nlargest(por_tipo, p).sort_values("hora") for s, p in tipos]
    return pd.concat(bloques).drop(columns="_norm")


def _autor(nombre) -> str:
    return re.sub(r"^@", "", nombre or "").strip()


def bloque(sel: pd.DataFrame, pagina: str, plataforma: str, link: str, resumen: str | None = None) -> str:
    lineas = [pagina.strip(), f" {plataforma.strip()} {link.strip()}".rstrip(), ""]
    if resumen:
        lineas += [resumen, ""]
    lineas += [f"{ICONOS[r.sentimiento]}{_autor(r.autor)} {r.mensaje.strip()}" for r in sel.itertuples()]
    return "\n".join(lineas)


def unir(bloques: list[str]) -> str:
    return "\n\n\n".join(bloques)


def resumen(df: pd.DataFrame) -> str:
    n = len(df)
    pct = df["sentimiento"].value_counts(normalize=True) * 100
    return (f"Comentarios analizados: {n:,} · 🟢 {pct.get('POS', 0):.0f}% positivos · "
            f"🔴 {pct.get('NEG', 0):.0f}% negativos · ⚪ {pct.get('NEU', 0):.0f}% neutros")


def word(contenido: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name, estilo.font.size = "Calibri", Pt(11)
    lineas = contenido.split("\n")
    for i, linea in enumerate(lineas):
        # El nombre de la página (inicio de cada bloque) va en negritas
        es_titulo = (i == 0 or not lineas[i - 1].strip()) and i + 1 < len(lineas) and lineas[i + 1].startswith(" ")
        p = doc.add_paragraph()
        p.add_run(linea).bold = es_titulo
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
